"""Generate the GeoSI presentation deck (GeoSI_Presentation.pptx)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

OUT = Path(__file__).resolve().parent.parent / "GeoSI_Presentation.pptx"

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
TEAL = RGBColor(0x10, 0x8A, 0x92)
SAND = RGBColor(0xF4, 0xF1, 0xEA)
INK = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x55, 0x5F, 0x6D)


def _band(slide, prs, color):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.35))
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def _title(slide, text, top=Inches(0.55), color=NAVY, size=36, bold=True):
    tb = slide.shapes.add_textbox(Inches(0.5), top, Inches(12), Inches(1.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.runs[0].font.size = Pt(size); p.runs[0].font.bold = bold; p.runs[0].font.color.rgb = color
    p.runs[0].font.name = "Calibri"


def _bullets(slide, items, top=Inches(1.7), left=Inches(0.6), width=Inches(12.0), height=Inches(5.2), size=18):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        for r in p.runs:
            r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = "Calibri"
        p.space_after = Pt(8)
        p.level = 0


def _footer(slide, prs, text="GeoSI  Geospatial Superintelligence  |  v2.0"):
    tb = slide.shapes.add_textbox(Inches(0.5), prs.slide_height - Inches(0.45), Inches(12), Inches(0.3))
    p = tb.text_frame.paragraphs[0]; p.text = text
    p.runs[0].font.size = Pt(10); p.runs[0].font.color.rgb = MUTED; p.runs[0].font.name = "Calibri"


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.2), prs.slide_width, Inches(0.06))
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(12), Inches(1.4))
    p = t.text_frame.paragraphs[0]; p.text = "GeoSI"
    p.runs[0].font.size = Pt(96); p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); p.runs[0].font.name = "Calibri"
    sub = s.shapes.add_textbox(Inches(0.8), Inches(4.35), Inches(12), Inches(0.7))
    p = sub.text_frame.paragraphs[0]; p.text = "The beginning of geospatial superintelligence."
    p.runs[0].font.size = Pt(26); p.runs[0].font.color.rgb = RGBColor(0xD9, 0xE6, 0xE9); p.runs[0].font.italic = True
    auth = s.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(12), Inches(0.5))
    p = auth.text_frame.paragraphs[0]; p.text = "Aaron R  Digital University Kerala  2026"
    p.runs[0].font.size = Pt(14); p.runs[0].font.color.rgb = RGBColor(0xA9, 0xBF, 0xC7)

    # 2. The Problem
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "The Problem")
    _bullets(s, [
        "Professional GIS is powerful but gated behind menus, scripts, and SQL.",
        "Analysts waste hours translating questions into toolchains.",
        "Domain experts cannot self-serve spatial insight.",
        "Existing assistants lack GIS context  they cannot run real geospatial tools.",
        "There is no conversational surface that plans and executes multi-step workflows on the layers you already have.",
    ])
    _footer(s, prs)

    # 3. The Vision
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "The Vision")
    _bullets(s, [
        "A conversational layer over professional GIS.",
        "Ask a spatial question  GeoSI plans a multi-step workflow and runs it.",
        "Local-first intelligence via Ollama; no API key required to start.",
        "Zero hallucinated data: GeoSI operates only on layers the user has loaded.",
        "Explainable: every result ships with the plan and the reasoning.",
    ])
    _footer(s, prs)

    # 4. Architecture
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Three-Surface Architecture")
    _bullets(s, [
        "geosi_engine/  universal core. No QGIS. No Qt. Importable anywhere.",
        "geosi_plugin/  the only code that touches QGIS. Reads the Layers panel.",
        "geosi_server/  optional FastAPI REST surface. OSM fallback when safe.",
        "Backend dispatch pattern keeps the engine framework-free.",
        "BackendTool  register_backend('qgis', fn) at plugin startup.",
    ])
    _footer(s, prs)

    # 5. Data Flow
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Data Flow: Prompt to Result")
    _bullets(s, [
        "1. User types a query in the chat dock.",
        "2. Plugin syncs QGIS Layers panel into engine state.",
        "3. Parser  AnalysisRequest (Ollama-first, rule-based fallback).",
        "4. Agent validates layers exist; returns clear error if not.",
        "5. Agent produces an ExecutionPlan (multi-step, explainable).",
        "6. Executor runs each ToolStep; BackendTool dispatches to QGIS Processing.",
        "7. Dock renders answer, reasoning, and suggestions.",
    ])
    _footer(s, prs)

    # 6. Capabilities
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Capabilities  125 Tools, 9 Domains")
    _bullets(s, [
        "Vector: 40 tools  buffer, intersect, clip, dissolve, voronoi, joins, CRS.",
        "Raster: 25 tools  NDVI/NDWI/NDBI, zonal stats, reclassify, resample.",
        "Terrain: 10 tools  slope, aspect, hillshade, contour, viewshed, watershed.",
        "Network: 8 tools  shortest path, service area, isochrone, OD matrix.",
        "AI/ML: 10 tools  K-means, DBSCAN, Getis-Ord Gi*, Moran's I, IDW.",
        "Cartography: 10 tools  GeoJSON, KML, Shapefile, geocode.",
        "Validation: 10 tools  fix geometries, topology, duplicates, CRS check.",
        "Temporal, portable I/O  auto-discovered via ToolRegistry.",
    ])
    _footer(s, prs)

    # 7. LLM chain
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "LLM Priority Chain")
    _bullets(s, [
        "1. Ollama (local)  always first. 1.5s probe, 30s generate timeout.",
        "2. Google Gemini  if GEMINI_API_KEY is set.",
        "3. Anthropic Claude  if ANTHROPIC_API_KEY is set.",
        "4. OpenAI  if OPENAI_API_KEY is set.",
        "5. Keyword rules  always on, drives all 125 tools offline.",
        "Every provider has fast failover; GeoSI never leaves the user hanging.",
    ])
    _footer(s, prs)

    # 8. Layer-not-found
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Grounded in Reality")
    _bullets(s, [
        "If a user references a layer that is not loaded, GeoSI refuses to guess.",
        "Fuzzy matcher finds near-matches and surfaces suggestions.",
        "Orphan-noun detector catches un-parsed references.",
        "Example: 'Buffer hospitals by 500m' with only Schools loaded  GeoSI replies: 'Layer hospitals not found. Did you mean Schools (50% match)?'",
        "Plugin never fetches data from OSM. Server fetches only when safe.",
    ])
    _footer(s, prs)

    # 9. Example workflows
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Example Workflows")
    _bullets(s, [
        "Buffer Schools by 1 km and intersect with TVM_Corp.",
        "Calculate NDVI from Landsat_red and Landsat_NIR.",
        "Cluster Assets using DBSCAN with eps=500.",
        "Shortest path from Schools to Assets over Roads.",
        "Reclassify DEM into low / medium / high elevation.",
        "Validate geometries in Kerala; fix and export as GeoJSON.",
    ])
    _footer(s, prs)

    # 10. API
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Public API (Python)")
    _bullets(s, [
        "from geosi_engine import GeoSI",
        "engine = GeoSI()               # auto-discovers 125 tools",
        "engine.state.add_layer(Layer(...))",
        "result = engine.query('Buffer Schools by 500m')",
        "result['success']              # bool",
        "result['answer']               # user-facing message",
        "result['plan']                 # list of ToolSteps",
        "result['reasoning']            # LLM or rule-based explanation",
    ])
    _footer(s, prs)

    # 11. Deployment
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Deployment")
    _bullets(s, [
        "Plugin: copy geosi_plugin/ + geosi_engine/ into QGIS plugin folder.",
        "Server: docker-ready FastAPI at geosi_server/ (optional).",
        "Ollama: local LLM, free, no API key. Pull mistral / llama3.1 / qwen2.5.",
        "Configuration via QgsSettings or environment variables.",
        "Zero cloud dependency for a full offline deployment.",
    ])
    _footer(s, prs)

    # 12. Roadmap
    s = prs.slides.add_slide(blank); _band(s, prs, TEAL); _title(s, "Roadmap")
    _bullets(s, [
        "v1.0  Core engine, 80 tools, basic plugin. (done)",
        "v2.0  Ollama-first, 125 tools, universal engine, validation. (done)",
        "v2.5  GeoPandas backend for server, batch mode, notebooks.",
        "v3.0  Real-time satellite streams, 3D viewshed, XR visualisation.",
        "v4.0  Predictive spatial modelling, autonomous decision support.",
    ])
    _footer(s, prs)

    # 13. Closing
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(12), Inches(1.2))
    p = t.text_frame.paragraphs[0]; p.text = "The beginning of geospatial superintelligence."
    p.runs[0].font.size = Pt(40); p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); p.runs[0].font.name = "Calibri"
    t2 = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(12), Inches(0.6))
    p = t2.text_frame.paragraphs[0]; p.text = "Aaron R  Digital University Kerala  aaronr.ds25@duk.ac.in"
    p.runs[0].font.size = Pt(18); p.runs[0].font.color.rgb = RGBColor(0xD9, 0xE6, 0xE9)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
